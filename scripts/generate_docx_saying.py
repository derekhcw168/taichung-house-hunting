import os
import sys

if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shading if any
    for shd in tcPr.findall(qn('w:shd')):
        tcPr.remove(shd)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    """Sets inner margins/padding for a cell (in dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for mar in tcPr.findall(qn('w:tcMar')):
        tcPr.remove(mar)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_left_border_only(cell, color_hex="1E3A8A", sz="24"):
    """Sets a solid left border and removes top, bottom, and right borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    for b in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(b)
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def set_cell_full_border(cell, color_hex="CBD5E1", sz="4"):
    """Sets subtle full border."""
    tcPr = cell._tc.get_or_add_tcPr()
    for b in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(b)
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def format_run(run, font_name="微軟正黑體", size_pt=10.5, bold=False, italic=False, color_rgb=(45, 55, 72)):
    """Applies font styles to run including East Asian font mapping."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)

def add_header_footer(doc):
    """Sets up page headers and footers."""
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("大墩人文 11 樓｜房仲向屋主議價說服專用說帖（內部實戰版）")
    format_run(hr, size_pt=8.5, color_rgb=(148, 163, 184))

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr1 = fp.add_run("機密資料・僅供經紀人內部洽談使用・請勿外流")
    format_run(fr1, size_pt=8.5, color_rgb=(148, 163, 184))
    
    # Add page number field to footer
    fp2 = footer.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr2 = fp2.add_run("第 ")
    format_run(fr2, size_pt=8.5, color_rgb=(148, 163, 184))
    
    fldSimple1 = OxmlElement('w:fldSimple')
    fldSimple1.set(qn('w:instr'), 'PAGE')
    fp2._p.append(fldSimple1)
    
    fr3 = fp2.add_run(" 頁")
    format_run(fr3, size_pt=8.5, color_rgb=(148, 163, 184))

def main():
    doc = docx.Document()
    add_header_footer(doc)

    # 1. Document Badge / Pre-title
    badge_p = doc.add_paragraph()
    badge_p.paragraph_format.space_before = Pt(0)
    badge_p.paragraph_format.space_after = Pt(4)
    run_badge = badge_p.add_run("【內部洽談專用】信義房屋專任委託物件・議價撮合策略說帖")
    format_run(run_badge, size_pt=9.5, bold=True, color_rgb=(37, 99, 235))

    # 2. Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("大墩人文 11 樓（專任捷運大三房）")
    format_run(run_title, size_pt=22, bold=True, color_rgb=(30, 58, 138))

    # 3. Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("房仲向屋主議價說服專用說帖（經紀人視角實戰版）")
    format_run(run_sub, size_pt=13, bold=True, color_rgb=(71, 85, 105))

    # 4. Meta Table (Quick Fact Sheet Box)
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("委託現況", "信義房屋專任約（屋齡 33.7 年）"),
        ("開價現狀", "950 萬元（單價約 30.12 萬/坪）"),
        ("買方出價", "斡旋 710 萬（上限 730 萬）"),
        ("買方人設", "不卑不亢理性自住家庭"),
        ("修繕扣除", "扣除基礎工程 120 萬元"),
        ("核心破口", "屋主繼承取得變現＋承諾善待起家厝")
    ]

    for idx, (label, val) in enumerate(meta_items):
        r_idx = idx // 3
        c_idx = idx % 3
        cell = meta_table.cell(r_idx, c_idx)
        cell.width = Inches(2.25)
        set_cell_background(cell, "F8FAFC")
        set_cell_full_border(cell, color_hex="E2E8F0", sz="4")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r_label = p.add_run(f"● {label}：")
        format_run(r_label, size_pt=9, bold=True, color_rgb=(71, 85, 105))
        r_val = p.add_run(val)
        format_run(r_val, size_pt=9, bold=False, color_rgb=(15, 23, 42))

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. Usage Note Callout Box
    note_table = doc.add_table(rows=1, cols=1)
    note_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    note_cell = note_table.cell(0, 0)
    note_cell.width = Inches(6.8)
    set_cell_background(note_cell, "EFF6FF")
    set_cell_left_border_only(note_cell, color_hex="2563EB", sz="28")
    set_cell_margins(note_cell, top=140, bottom=140, left=200, right=200)

    np = note_cell.paragraphs[0]
    np.paragraph_format.space_after = Pt(4)
    nr_title = np.add_run("📋 本說帖使用說明與核心戰略方針：")
    format_run(nr_title, size_pt=10.5, bold=True, color_rgb=(30, 58, 138))

    points = [
        ("經紀人溝通視角", "全篇以「站在屋主利益角度著想、剖析客觀市場現實、突顯稀缺買方誠意」為出發點。將買方抗性轉化為房仲替屋主分析局勢、軟化價格防線的專業對談。"),
        ("買方人設定位", "不卑不亢、理性務實、財務紀律嚴謹的優質自住家庭。既不哭窮裝可憐（避免屋主疑慮違約），也不吹噓揮霍（避免成為被宰肥羊）。"),
        ("雙核心切入點", "① 點出屋主為「繼承取得」，原始持有成本極低，及早變現即是全額純獲利落袋；② 買方承諾投入 120 萬翻新水電泥作，「用心善待長輩起家厝」，達成安心託付的圓滿善緣。")
    ]
    for p_title, p_desc in points:
        p_pt = note_cell.add_paragraph()
        p_pt.paragraph_format.space_after = Pt(3)
        p_pt.paragraph_format.line_spacing = 1.2
        r1 = p_pt.add_run(f"• {p_title}：")
        format_run(r1, size_pt=9.5, bold=True, color_rgb=(30, 58, 138))
        r2 = p_pt.add_run(p_desc)
        format_run(r2, size_pt=9.5, bold=False, color_rgb=(51, 65, 85))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 6. Section 1: 開場破冰與局勢鋪陳
    h1_1 = doc.add_paragraph()
    h1_1.paragraph_format.space_before = Pt(14)
    h1_1.paragraph_format.space_after = Pt(6)
    r_h1_1 = h1_1.add_run("【第一階段：開場破冰與局勢鋪陳】")
    format_run(r_h1_1, size_pt=14, bold=True, color_rgb=(30, 58, 138))

    sub_h1 = doc.add_paragraph()
    sub_h1.paragraph_format.space_after = Pt(6)
    r_sub_h1 = sub_h1.add_run("🎯 房仲開場白（站在屋主立場，傳遞市場寒意、買方理性預算與真實誠意）")
    format_run(r_sub_h1, size_pt=11, bold=True, color_rgb=(37, 99, 235))

    # Pitch Box 1
    pitch_1_text = (
        "「大姐您好，今天特地跟您回報這陣子的帶看進度與市場現況。\n\n"
        "相信您最近看新聞也知道，自從央行第七波信用管制實施、銀行房貸滿水位排隊之後，整個台中市的買氣直接急凍了三到四成。現在市場上出來看屋的客人大多只是觀望，而且很多都是投資客隨便亂出芭樂價。\n\n"
        "但今天我要跟您回報一個好消息：我們店頭遇到了一組非常優質、理性且誠意十足的自住家庭！\n\n"
        "這組客人做足了市場功課，非常喜歡大姐這間房子的生活圈與格局，態度非常務實。他們有嚴謹的家庭財務規劃與預算紀律，看完之後正式委託我們表達意願。\n\n"
        "坦白跟大姐報告，買方態度不卑不亢，對價格有非常客觀且清晰的計算：他們評估了 34 年老屋必做的全室水電與廚衛基礎工程，將房屋購買價嚴格設定在合理的財務上限。我是站在協助大姐『安全變現獲利、落袋為安』，同時也讓長輩傳承下來的好房子找到珍惜它的新主人的立場，把市場客觀現實與買方的理性評估向您老實分析，我們一起來看看怎麼圓滿促成！」"
    )

    p_table1 = doc.add_table(rows=1, cols=1)
    p_table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1 = p_table1.cell(0, 0)
    c1.width = Inches(6.8)
    set_cell_background(c1, "F8FAFC")
    set_cell_left_border_only(c1, color_hex="1E3A8A", sz="28")
    set_cell_margins(c1, top=140, bottom=140, left=200, right=200)

    p_c1 = c1.paragraphs[0]
    p_c1.paragraph_format.line_spacing = 1.3
    p_c1.paragraph_format.space_after = Pt(0)
    r_pitch1 = p_c1.add_run(pitch_1_text)
    format_run(r_pitch1, size_pt=10, bold=False, color_rgb=(30, 41, 59))

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 7. Section 2: 12 大核心理由與話術逐字稿
    h1_2 = doc.add_paragraph()
    h1_2.paragraph_format.space_before = Pt(16)
    h1_2.paragraph_format.space_after = Pt(8)
    r_h1_2 = h1_2.add_run("【第二階段：房仲說服屋主的 12 大核心理由與話術逐字稿】")
    format_run(r_h1_2, size_pt=14, bold=True, color_rgb=(30, 58, 138))

    reasons = [
        {
            "num": 1,
            "title": "【市場大環境急凍與理性自住客的稀缺性】做足功課且隨時能履約的優質自住家庭",
            "logic": "營造錯過當前做足功課、財務健全、隨時能履約之理性自住買方的沈沒成本與時間風險。",
            "script": "「大姐，現在房貸利率實質都拉到近 3% 了，市場買氣直接凍結。現在市場上很多看屋客不是觀望就是亂喊芭樂價。這位買方是我們店頭難得遇到『做足功課、財務規劃健全、隨時能開立斡旋支票簽約』的理性自住家庭。買方態度很誠懇，但也非常講求財務紀律。以目前市場這麼凍的狀況，如果我們這次價格踩得太硬把這組理性客人放走，房子繼續掛半年，下一組願意出價的買方不知道要等幾個月，到時候市場可能又更冷了。」"
        },
        {
            "num": 2,
            "title": "【實登天花板已現與長期均價分析】開價 950 萬已在歷史頂峰，遠超長期每坪 20 萬合理水位",
            "logic": "客觀指出 30 萬/坪是社區歷史頂峰，且對比 2015～2022 年長達 7~8 年歷史均價僅約 20 萬/坪，點出目前價格已大幅偏離長期合理基準，買方無追價意願。",
            "script": "「大姐，我們開價從 998 萬調整到 950 萬，折合單價 30.12 萬/坪，雖然對外有降價動作，但現在買方每個人手機都滑實價登錄，精得很。\n\n他們查得很清楚：這個社區 34 年來歷史最高單價也就 30.3 萬，買方會覺得我們開價還是踩在歷史最頂峰。\n\n更關鍵的是，買方調出實價登錄歷史走勢，2015 年到 2022 年中間長達七、八年的時間，我們社區的長期合理行情一直都維持在一坪 20 萬元左右！現在開價 30 萬以上，已經比當初長期行情的合理價格高出非常多（漲了整整五成以上）！\n\n現在大環境景氣反轉、央行打炒房，完全沒有買方願意再追高去買老國宅，買方出價只是希望能回歸理性合理的水位，我們適度讓利才能促成共識。」"
        },
        {
            "num": 3,
            "title": "【無獨立產權車位】南屯三房無固定車位之市場硬傷",
            "logic": "普遍三房買方必備車位，每年抽籤的抗性在中古市場有實質折價。",
            "script": "「南屯市區這邊買三房的家庭，九成以上都有開車，對固定車位需求很高。我們社區車位產權是公有的、每年要碰運氣抽籤，沒抽到就得去外面吹風淋雨找月租，這點在帶看時經常被其他組客人打槍。現在這組買方願意接受抽籤條件已經很難得，但他堅持沒有產權車位的大樓，單價本來就應該比周邊附車位的大樓親民一成，這點在行情比價上我們確實比較難反駁他。」"
        },
        {
            "num": 4,
            "title": "【買方工作在霧峰，隨時可能轉移陣地】感情牌 vs 理智比價",
            "logic": "點出買方並非「非買南屯不可」，南區距離工作地點更近，他手上隨時有其他物件可選。",
            "script": "「這組買方其實對我們社區很有感情，因為他之前在這裡租過房子，喜歡這裡的生活圈，所以我才能一路把他拉著談到現在。但他實際上是在霧峰上班，他同時也在看南區那邊的大樓（南區總價 800 多萬，室內更大、還附汽車車位）。如果我們價格卡太死，他理智一回神就會直接轉去買南區了，等於我們白白流失一個最有感情基礎的買方。」"
        },
        {
            "num": 5,
            "title": "【全室老屋基礎工程剛性必要】120 萬水電泥作修繕是客觀成本，買方理性扣除",
            "logic": "點出 34 年老屋必做水電管線與衛浴更新，120 萬是不可避免的剛性安全支出，買方理性納入總成本扣除，非惡意壓價。",
            "script": "「現在工錢、泥作跟垃圾清運費漲得非常誇張。買方看完後非常客觀地理性評估：戶內原本的舊櫥櫃和固定裝潢全都要拆除清運，加上快 34 年的老舊電線、冷熱水管、兩間衛浴跟廚具全部必須打掉重練，統包報價的基本安全工程就要 100 到 120 萬。\n\n買方態度很實在，這筆費用是為了長久自住安全必須投入的剛性成本，他們規劃以裝潢修繕貸款搭配自有資金支應。在理性財務規劃下，這筆修繕款自然必須從房屋購買總價中扣除，這完全是基於老屋現況的客觀計算，而不是故意挑毛病殺價。」"
        },
        {
            "num": 6,
            "title": "【屋齡逼近 34 年】趁首購還能貸款時順利出脫，避免越老越難賣",
            "logic": "提醒屋主未來資產折舊與貸款成數縮水的時間成本，趁老屋還能貸款時落袋為安。",
            "script": "「大姐，我們這棟屋齡已經快 34 年了。在銀行貸款實務上，屋齡超過 30 年審核就會逐步變嚴，再過幾年屋齡逼近 40 年時，能接手的買方客群會越來越窄，房貸成數跟年限都會被壓縮。趁現在屋齡還在一般首購族願意承接的最後黃金期順利出脫、獲利了結，絕對比再放幾年越來越老更好處理。」"
        },
        {
            "num": 7,
            "title": "【銀行鑑價客觀防線】無車位老國宅鑑價保守，買方嚴守財務紀律不盲目追高",
            "logic": "搬出金融機構客觀估價防線，買方重視家庭財務紀律，絕不冒著鑑價落差風險盲目過度追高。",
            "script": "「現在銀行房貸審核極度嚴格，鑑價公司在估 34 年、無車位的集合國宅時態度都很保守，行員私下評估這間很難鑑價到 950 萬的天花板。買方是一組重視財務紀律的理性家庭，他們願意按常態貸款購屋，但絕不可能去冒著銀行鑑價落差的風險盲目硬追高，這也是他們出價踩在客觀合理區間的原因。」"
        },
        {
            "num": 8,
            "title": "【實登近期低點定錨】114 年 10 月 2 樓成交 800 萬的市場訊號",
            "logic": "拿最新成交的 800 萬打擊屋主高價預期，證明市場價格已開始鬆動。",
            "script": "「買方現在出價都拿最新的成交紀錄跟我們對。像 114 年 10 月我們社區才剛成交一戶 800 萬（單價約 25.5 萬/坪）。買方現在心態被這筆 800 萬定錨住，一直跟我們爭取說低樓層都能成交到 800 萬，現在買氣又比那時候更冷，他認為整棟合理的成交重心本來就應該在 800 萬之下。」"
        },
        {
            "num": 9,
            "title": "【買方家庭預算客觀精算】850 萬總配置扣除 120 萬修繕，購買極限天花板 730 萬（出 710 萬保留合理磋商空間）",
            "logic": "將買方的出價轉化為「透明理性的家庭財務預算配置」，證明 710 萬斡旋是兼顧屋況與磋商空間的務實出價。",
            "script": "「大姐，我老老實實把這組自住買方的財務算盤攤給您看：他們家庭經過審慎評估，買這間房子的全部資金預算上限設定在 850 萬。\n\n但這間 34 年老屋，拆除清運、重拉全室水電、兩間衛浴與廚具重做、氣密窗與壁癌泥作，最精準陽春的基礎工程報價單就要 120 萬（這筆預計搭配裝潢修繕貸款規劃）。\n\n850 萬總預算扣掉 120 萬工程款，買方買這間房子的購買極限天花板就是 730 萬。\n\n他今天之所以出 710 萬下斡旋，完全是展現誠意、並保留 20 萬空間給我們業務協助雙方圓滿撮合在 730 萬。這完全是一套不卑不亢、有依有據的理性預算配置，絕非隨便胡亂喊價。」"
        },
        {
            "num": 10,
            "title": "【周邊中古競品比價效應】同價位區間，周邊有更大或附車位物件",
            "logic": "點出市場上的同價位替代品，證明無車位在市場上缺乏競爭力。",
            "script": "「我們如果在市場上跟周邊比，黎明商圈或南區，相同價位區間的中古大樓，很多都已經能買到附車位的物件。大墩人文開 950 萬、而且還沒車位，在市場比價性上確實比較吃虧。除非價格能讓買方覺得撿到便宜，否則自住客很容易就被周邊有車位的競品拉走了。」"
        },
        {
            "num": 11,
            "title": "【強烈具體競品施壓與最後通牒】永義帶看「靜園」水電已換、美裝潢免花 120 萬可直入，大墩不談買方立刻轉移！",
            "logic": "點名實體同業競爭對手（永義房屋 vs 具體社區「靜園」），營造隨時會被別家房仲整碗捧走、買方轉身就走的致命急迫感（FOMO），打破屋主「慢慢拖慢慢賣」的僥倖心態。",
            "script": "「大姐，我跟您透露一個非常緊急的內部消息：這組買方除了我們這間，這兩天永義房屋的業務也在積極帶他看南區的『靜園』社區。\n\n買方看完『靜園』也非常心動，因為靜園那一戶原屋主已經把全室基礎工程、水電管線全部重新做過，而且還做了非常漂亮的高質感裝潢，買方買下來完全不用再多掏 120 萬基礎工程費用，直接搬行李就能入住！\n\n買方今天跟我講得很明白：他雖然對大墩人文有感情，但如果大墩人文大姐您這邊堅持不願意讓利、沒有意願要談的話，他今天回去就會準備去跟永義房屋談下一個『靜園』這個物件了！\n\n大姐，我們現在是跟別家房仲在搶這組客人，要是今天我們猶豫放手，買方一轉頭跟永義簽約，我們就真的永遠失去這個好買方了，請大姐一定要把握今天最後的成交機會！」"
        },
        {
            "num": 12,
            "title": "【屋主繼承取得成本極低，安全變現落袋為安＋買方承諾用心善待起家厝】",
            "logic": "點出屋主為繼承取得，原始持有成本極低，變現即是全額純獲利；買方承諾投入上百萬細心翻修，好好善待房子，兼顧屋主獲利與老屋善緣交託。",
            "script": "「大姐，其實我們私底下幫您算過，您這間房子是因為長輩傳承、繼承取得的，原始持有成本非常低，基本上沒有什麼沉重的房貸負擔。今天只要能順利成交變現，不管是七百多萬的哪一個數字，對大姐來說全額都是實質大筆獲利落袋，把不動產轉換為最安全靈活的現金，更省去後續每年繳房屋稅、地價稅以及大樓管理的煩惱！\n\n而且最難得的是，這組買方態度非常好，他們不是投資客買來隨便隔套出租，而是打算一家人長期在這裡安居。他們甚至願意自己另外規劃 120 萬的修繕工程，把全室水電管線、泥作防水全部重新做好。買方親口跟我保證：『一定會好好珍惜、用心對待大姐長輩傳承下來的這個起家厝！』\n\n大姐，把房子賣給這樣有心、有原則、會好好善待屋子的優質自住家庭，既讓大姐的大筆資產順利變現落袋，也為老房子找到最好的託付與傳承，這真的是最圓滿的善緣！」"
        }
    ]

    for item in reasons:
        # Title of Reason
        rp = doc.add_paragraph()
        rp.paragraph_format.space_before = Pt(11)
        rp.paragraph_format.space_after = Pt(2)
        r_num_badge = rp.add_run(f"理由 {item['num']:02d}｜")
        format_run(r_num_badge, size_pt=11.5, bold=True, color_rgb=(37, 99, 235))
        r_item_title = rp.add_run(item['title'])
        format_run(r_item_title, size_pt=11.5, bold=True, color_rgb=(30, 58, 138))

        # Logic / Strategy
        lp = doc.add_paragraph()
        lp.paragraph_format.space_before = Pt(0)
        lp.paragraph_format.space_after = Pt(4)
        r_logic_tag = lp.add_run("💡 說服核心邏輯：")
        format_run(r_logic_tag, size_pt=9.5, bold=True, color_rgb=(100, 116, 139))
        r_logic_val = lp.add_run(item['logic'])
        format_run(r_logic_val, size_pt=9.5, bold=False, italic=False, color_rgb=(71, 85, 105))

        # Agent Pitch Callout Box
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.width = Inches(6.8)
        
        # Color accent: special border for reasons 9, 11, 12 (key attack points)
        if item['num'] in [9, 11, 12]:
            set_cell_background(cell, "F0FDF4" if item['num'] == 12 else "FEF2F2" if item['num'] == 11 else "EFF6FF")
            border_color = "16A34A" if item['num'] == 12 else "DC2626" if item['num'] == 11 else "2563EB"
            set_cell_left_border_only(cell, color_hex=border_color, sz="28")
        else:
            set_cell_background(cell, "F8FAFC")
            set_cell_left_border_only(cell, color_hex="3B82F6", sz="20")

            
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

        cp = cell.paragraphs[0]
        cp.paragraph_format.line_spacing = 1.25
        cp.paragraph_format.space_after = Pt(0)
        
        r_prefix = cp.add_run("🗣️ 經紀人向屋主說法：\n")
        format_run(r_prefix, size_pt=9.5, bold=True, color_rgb=(30, 58, 138))
        
        r_script = cp.add_run(item['script'])
        format_run(r_script, size_pt=10, bold=False, color_rgb=(30, 41, 59))

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 8. Section 3: 收尾逼單促成簽字
    h1_3 = doc.add_paragraph()
    h1_3.paragraph_format.space_before = Pt(16)
    h1_3.paragraph_format.space_after = Pt(6)
    r_h1_3 = h1_3.add_run("【第三階段：收尾逼單促成簽字】")
    format_run(r_h1_3, size_pt=14, bold=True, color_rgb=(30, 58, 138))

    sub_h3 = doc.add_paragraph()
    sub_h3.paragraph_format.space_after = Pt(6)
    r_sub_h3 = sub_h3.add_run("🏆 拋出「理性自住家庭 10 萬真金白銀斡旋金，不卑不亢促成買賣雙贏」")
    format_run(r_sub_h3, size_pt=11, bold=True, color_rgb=(217, 119, 6))

    # Closing Callout Box
    c_table = doc.add_table(rows=1, cols=1)
    c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cc = c_table.cell(0, 0)
    cc.width = Inches(6.8)
    set_cell_background(cc, "FFFBEB")
    set_cell_left_border_only(cc, color_hex="D97706", sz="32")
    set_cell_margins(cc, top=140, bottom=140, left=200, right=200)

    p_cc = cc.paragraphs[0]
    p_cc.paragraph_format.line_spacing = 1.3
    p_cc.paragraph_format.space_after = Pt(0)
    
    r_cc_head = p_cc.add_run("🎯 經紀人逼單收尾最關鍵的一句話（激發屋主落袋為安與善緣交託心態）：\n\n")
    format_run(r_cc_head, size_pt=10.5, bold=True, color_rgb=(180, 83, 9))

    closing_text = (
        "「大姐，現在市場上很多買方只是隨便出芭樂價問問、根本沒有誠意買。但這組買方態度不卑不亢、理性堅定，直接備妥 10 萬真金白銀的斡旋支票放在我們店頭桌上，展現隨時能簽約履約的十足誠意！\n\n"
        "買方出價有理有據，評估客觀，既不刻意挑剔打壓，也不盲目過度追價。\n\n"
        "大姐是繼承取得，成本極低，能在大環境反轉前及早變現落袋才是真金白銀。把房子交給這組會花心思重新整頓水電、承諾會好好善待起家厝的優質自住家庭，既圓滿了資產處分，也圓滿了一段美好的傳承。請大姐今天就成全，我們直接簽字圓滿成交！」"
    )
    r_closing = p_cc.add_run(closing_text)
    format_run(r_closing, size_pt=10.5, bold=False, color_rgb=(30, 41, 59))

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 9. Section 4: 房仲心理戰術技巧提醒
    h1_4 = doc.add_paragraph()
    h1_4.paragraph_format.space_before = Pt(14)
    h1_4.paragraph_format.space_after = Pt(6)
    r_h1_4 = h1_4.add_run("【總結備忘：經紀人實戰心理攻防守則】")
    format_run(r_h1_4, size_pt=13, bold=True, color_rgb=(30, 58, 138))

    tips = [
        ("循序漸進，不要一次全丟", "先從宏觀環境急凍與無車位抗性切入軟化屋主防線；第二階段丟出 120 萬基礎工程客觀支出與競品『靜園』緊迫通牒；進入收網階段再打出『屋主繼承取得全額變現』與『買方用心善待起家厝』，最後亮出 10 萬真金白銀支票逼定。"),
        ("給足經紀人談判底氣", "買方下斡旋時務必備妥 10 萬支票或現金，讓經紀人面對專任委託屋主時能拍胸脯保證：『這是隨時能簽約的真金白銀，錯過這組可能要再等半年』，徹底擊潰屋主惜售拖延的僥倖心態。")
    ]

    for t_title, t_desc in tips:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(2)
        tp.paragraph_format.space_after = Pt(4)
        tp.paragraph_format.line_spacing = 1.2
        r_tt = tp.add_run(f"✔ {t_title}：")
        format_run(r_tt, size_pt=10, bold=True, color_rgb=(30, 58, 138))
        r_td = tp.add_run(t_desc)
        format_run(r_td, size_pt=10, bold=False, color_rgb=(51, 65, 85))

    # Output path
    output_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "大墩人文", "大墩人文11樓_房仲向屋主議價說服專用說帖.docx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"✅ 成功生成 Word 檔案：{output_path}")

if __name__ == "__main__":
    main()
